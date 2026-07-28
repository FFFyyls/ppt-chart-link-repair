#!/usr/bin/env python3
# Copyright 2026 XLoffice-Fyl
# SPDX-License-Identifier: Apache-2.0

"""Prepare native-link recipes and remove temporary Excel carrier sheets."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from native_link_pipeline import build_chart_recipes


def load_json(path: str | Path) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_json(path: str | Path, value: Any) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def prepare(manifest_path: str | Path, output_path: str | Path) -> dict[str, Any]:
    manifest = load_json(manifest_path)
    recipes = build_chart_recipes(manifest)
    result = {
        "schema_version": "2.0",
        "source_pptx": manifest.get("source_pptx"),
        "output_pptx": manifest.get("output_pptx"),
        "work_dir": manifest.get("work_dir"),
        "charts": recipes,
        "groups": manifest.get("groups", []),
    }
    write_json(output_path, result)
    return result


def clean_support_sheets(
    recipe_manifest_path: str | Path, support_map_path: str | Path
) -> dict[str, Any]:
    recipes = load_json(recipe_manifest_path)
    support = json.loads(Path(support_map_path).read_text(encoding="utf-8"))
    by_workbook: dict[str, set[str]] = {}
    for entry in support:
        by_workbook.setdefault(str(Path(entry["workbook_path"]).resolve()), set()).add(
            str(entry["support_sheet"])
        )
    cleaned: list[dict[str, Any]] = []
    for workbook_path, sheet_names in by_workbook.items():
        workbook = load_workbook(workbook_path)
        try:
            removed: list[str] = []
            for sheet_name in sorted(sheet_names):
                if sheet_name in workbook.sheetnames and len(workbook.sheetnames) > 1:
                    del workbook[sheet_name]
                    removed.append(sheet_name)
            workbook.save(workbook_path)
        finally:
            workbook.close()
        reopened = load_workbook(workbook_path, read_only=True, data_only=False)
        try:
            remaining = sorted(set(sheet_names).intersection(reopened.sheetnames))
            if remaining:
                raise RuntimeError(
                    f"Temporary carrier sheets remain in {workbook_path}: {remaining}"
                )
        finally:
            reopened.close()
        cleaned.append({"workbook_path": workbook_path, "removed": removed})
    return {"passed": True, "workbooks": cleaned, "chart_count": len(recipes["charts"])}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    prepare_parser = subparsers.add_parser("prepare")
    prepare_parser.add_argument("--manifest", required=True)
    prepare_parser.add_argument("--output", required=True)
    clean_parser = subparsers.add_parser("clean")
    clean_parser.add_argument("--recipes", required=True)
    clean_parser.add_argument("--support-map", required=True)
    clean_parser.add_argument("--result", required=True)
    args = parser.parse_args()
    if args.command == "prepare":
        result = prepare(args.manifest, args.output)
        print(json.dumps({"charts": len(result["charts"])}, ensure_ascii=False))
        return 0
    result = clean_support_sheets(args.recipes, args.support_map)
    write_json(args.result, result)
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
