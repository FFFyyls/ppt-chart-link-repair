# Copyright 2026 XLoffice-Fyl
# SPDX-License-Identifier: Apache-2.0

import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

import build_carrier_workbooks as carriers
import chart_recovery as cr
import office_verify_manifest as office_manifest
import portable_bundle


def make_fixture(
    path: Path,
    sparse: bool = False,
    inline_arrays: bool = False,
    external_target: str = "E:/ProjectA/data.xlsx",
) -> None:
    if sparse:
        cat_points = '<c:pt idx="0"><c:v>2024</c:v></c:pt><c:pt idx="2"><c:v>2026</c:v></c:pt>'
        value_points = '<c:pt idx="0"><c:v>10</c:v></c:pt><c:pt idx="2"><c:v>30</c:v></c:pt>'
        range_end = "D"
        point_count = 3
    else:
        cat_points = '<c:pt idx="0"><c:v>2024</c:v></c:pt><c:pt idx="1"><c:v>2025</c:v></c:pt>'
        value_points = '<c:pt idx="0"><c:v>10</c:v></c:pt><c:pt idx="1"><c:v>20</c:v></c:pt>'
        range_end = "C"
        point_count = 2
    if inline_arrays:
        tx_formula = "sheet"
        cat_formula = '{"2024","2025"}'
        val_formula = "{10,20}"
    else:
        tx_formula = "'[data.xlsx]Sheet A'!$A$2"
        cat_formula = f"'[data.xlsx]Sheet A'!$B$1:${range_end}$1"
        val_formula = f"'[data.xlsx]Sheet A'!$B$2:${range_end}$2"
    chart = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
 <c:chart><c:plotArea><c:barChart><c:barDir val="col"/><c:grouping val="clustered"/><c:ser>
  <c:idx val="0"/><c:order val="0"/>
  <c:tx><c:strRef><c:f>{tx_formula}</c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>Series 1</c:v></c:pt></c:strCache></c:strRef></c:tx>
  <c:cat><c:strRef><c:f>{cat_formula}</c:f><c:strCache><c:ptCount val="{point_count}"/>{cat_points}</c:strCache></c:strRef></c:cat>
  <c:val><c:numRef><c:f>{val_formula}</c:f><c:numCache><c:formatCode>0.00</c:formatCode><c:ptCount val="{point_count}"/>{value_points}</c:numCache></c:numRef></c:val>
 </c:ser></c:barChart></c:plotArea></c:chart>
 <c:externalData r:id="rId1"><c:autoUpdate val="0"/></c:externalData>
</c:chartSpace>'''
    slide = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
 <p:cSld><p:spTree>
  <p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="2" name="Chart 1"/></p:nvGraphicFramePr><a:graphic><a:graphicData><c:chart r:id="rId1"/></a:graphicData></a:graphic></p:graphicFrame>
  <p:sp><p:txBody><a:p><a:r><a:t>2020–2025年1–11月</a:t></a:r></a:p></p:txBody></p:sp>
  <p:sp><p:txBody><a:p><a:r><a:t>1</a:t></a:r></a:p></p:txBody></p:sp>
 </p:spTree></p:cSld>
</p:sld>'''
    slide_rels = '''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
 <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="../charts/chart1.xml"/>
</Relationships>'''
    chart_rels = '''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
 <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="{external_target}" TargetMode="External"/>
</Relationships>'''
    content_types = '''<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
 <Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
</Types>'''
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("ppt/slides/slide1.xml", slide)
        archive.writestr("ppt/slides/_rels/slide1.xml.rels", slide_rels)
        archive.writestr("ppt/charts/chart1.xml", chart)
        archive.writestr("ppt/charts/_rels/chart1.xml.rels", chart_rels)
        archive.writestr("ppt/charts/style1.xml", b"<style/>")
        archive.writestr("ppt/charts/colors1.xml", b"<colors/>")


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def _recover(self, source: Path, name: str = "work") -> tuple[dict, Path]:
        inventory = cr.inspect_pptx(source, self.root / name)
        inventory["selection_confirmed"] = True
        inventory["charts"][0]["selected"] = True
        inventory_path = Path(inventory["inventory_path"])
        inventory_path.write_text(
            json.dumps(inventory, ensure_ascii=False), encoding="utf-8"
        )
        repair = cr.recover_manifest(inventory_path)
        return repair, Path(repair["manifest_path"])

    def test_inspect_classifies_complete_and_partial_caches(self):
        complete = self.root / "complete.pptx"
        partial = self.root / "partial.pptx"
        make_fixture(complete)
        make_fixture(partial, sparse=True)
        complete_inventory = cr.inspect_pptx(complete, self.root / "complete-work")
        partial_inventory = cr.inspect_pptx(partial, self.root / "partial-work")
        self.assertEqual(complete_inventory["charts"][0]["recoverability"], "A")
        self.assertEqual(partial_inventory["charts"][0]["recoverability"], "B")

    def test_inline_array_formulas_materialize_as_workbook_ranges(self):
        source = self.root / "inline-arrays.pptx"
        make_fixture(source, inline_arrays=True, external_target="NULL")

        repair, _ = self._recover(source)

        self.assertTrue(repair["relink_allowed"])
        mapping = repair["charts"][0]["data_mappings"][0]
        self.assertEqual(["A2", "A3"], mapping["category_cells"])
        self.assertEqual(["B2", "B3"], mapping["value_cells"])
        workbook = Path(repair["groups"][0]["workbook_path"])
        opened = load_workbook(workbook, read_only=True, data_only=False)
        try:
            sheet = opened[mapping["sheet"]]
            self.assertEqual([10.0, 20.0], [sheet["B2"].value, sheet["B3"].value])
        finally:
            opened.close()
        recipes = carriers.build_chart_recipes(repair)
        refs = recipes[0]["series"][0]["refs"]
        self.assertEqual("range", refs["cat"]["parsed"]["kind"])
        self.assertEqual("range", refs["val"]["parsed"]["kind"])

    def test_confirm_selection_sets_only_explicit_chart_ids(self):
        source = self.root / "selection.pptx"
        make_fixture(source)
        inventory = cr.inspect_pptx(source, self.root / "selection-work")
        confirmed = cr.confirm_selection(
            inventory["inventory_path"], ["slide-1-chart-1"]
        )
        self.assertTrue(confirmed["selection_confirmed"])
        self.assertEqual(confirmed["confirmed_chart_ids"], ["slide-1-chart-1"])
        self.assertTrue(confirmed["charts"][0]["selected"])

    def test_recover_and_apply_update_change_only_rebuilt_workbook(self):
        source = self.root / "source.pptx"
        make_fixture(source)
        source_hash = cr.sha256_file(source)
        repair, manifest_path = self._recover(source)
        workbook = Path(repair["groups"][0]["workbook_path"])
        updates = self.root / "updates.json"
        updates.write_text(
            json.dumps(
                {
                    "updates_confirmed": True,
                    "updates": [
                        {
                            "chart_id": "slide-1-chart-1",
                            "sheet": "Sheet A",
                            "target_cell": "C2",
                            "old_value": 20.0,
                            "new_value": 25.0,
                            "source_note": "unit test",
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )
        result = cr.apply_updates(manifest_path, updates)
        book = load_workbook(workbook, data_only=False)
        try:
            self.assertEqual(book["Sheet A"]["C2"].value, 25.0)
        finally:
            book.close()
        self.assertTrue(result["relink_allowed"])
        self.assertEqual(cr.sha256_file(source), source_hash)

    def test_recover_can_stage_all_delivery_paths_before_office_work(self):
        source = self.root / "delivery-source.pptx"
        make_fixture(source)
        inventory = cr.inspect_pptx(source, self.root / "delivery-work")
        inventory["selection_confirmed"] = True
        inventory["charts"][0]["selected"] = True
        inventory_path = Path(inventory["inventory_path"])
        inventory_path.write_text(
            json.dumps(inventory, ensure_ascii=False), encoding="utf-8"
        )
        delivery = (self.root / "final-delivery").resolve()

        repair = cr.recover_manifest(inventory_path, delivery)

        self.assertEqual(
            Path(repair["groups"][0]["workbook_path"]).parent,
            delivery / "recovered-workbooks",
        )
        self.assertEqual(
            Path(repair["output_pptx"]), delivery / "delivery-source-relinked.pptx"
        )
        self.assertEqual(
            Path(repair["deliverable_report"]), delivery / "数据更新修复报告.md"
        )
        args = cr.build_parser().parse_args(
            ["recover", "--manifest", str(inventory_path), "--output-dir", str(delivery)]
        )
        self.assertEqual(Path(args.output_dir), delivery)

        report = cr._write_report(repair, [], None)
        self.assertEqual(report, delivery / "数据更新修复报告.md")
        self.assertTrue(report.exists())

    def test_partial_cache_creates_draft_but_blocks_relink(self):
        source = self.root / "partial.pptx"
        make_fixture(source, sparse=True)
        repair, manifest_path = self._recover(source, "partial-work")
        self.assertFalse(repair["relink_allowed"])
        self.assertTrue(Path(repair["groups"][0]["workbook_path"]).exists())
        with self.assertRaises(cr.RecoveryBlocked):
            cr.relink_manifest(manifest_path)

    def test_native_engine_contract_replaces_obsolete_relink(self):
        production = "\n".join(
            path.read_text(encoding="utf-8-sig")
            for pattern in ("*.py", "*.ps1")
            for path in SCRIPTS.glob(pattern)
            if path.name != "quick_validate.py"
        )
        self.assertNotIn("Chart." + "SetSourceData", production)
        self.assertNotIn("linked Excel chart object " + "fallback", production)
        self.assertIn("PasteSourceFormatting", production)
        self.assertIn("transplant_native_links.py", production)

    def test_office_verifier_has_gate_fresh_reopen_and_mutation(self):
        verifier = (SCRIPTS / "office_verify.ps1").read_text(encoding="utf-8-sig")
        worker = (SCRIPTS / "office_stage_worker.ps1").read_text(encoding="utf-8-sig")
        self.assertIn("PowerPoint正在运行，请保存并关闭所有PowerPoint窗口后继续", verifier)
        self.assertIn("exit 20", verifier)
        self.assertIn("EditData", worker)
        self.assertIn("WorkbookPath", worker)
        self.assertIn("MutationUpdate", worker)
        self.assertIn("ReadValue", worker)
        self.assertIn("user_excel_processes_preserved", verifier)

    def test_office_verifier_refuses_before_chartdata_when_excel_is_running(self):
        verifier = (SCRIPTS / "office_verify.ps1").read_text(encoding="utf-8-sig")
        excel_gate = verifier.index("$excelPids = @(Get-Process EXCEL")
        manifest_build = verifier.index("& $PythonPath $builder")

        self.assertLess(excel_gate, manifest_build)
        self.assertIn("Excel正在运行", verifier[excel_gate:manifest_build])
        self.assertIn("exit 20", verifier[excel_gate:manifest_build])

    def test_office_verifier_combines_access_checks_but_keeps_update_fresh(self):
        verifier = (SCRIPTS / "office_verify.ps1").read_text(encoding="utf-8-sig")
        worker = (SCRIPTS / "office_stage_worker.ps1").read_text(encoding="utf-8-sig")
        self.assertIn("Invoke-Worker 'Primary'", verifier)
        self.assertIn("Invoke-Worker 'ReadValue'", verifier)
        self.assertNotIn("Invoke-Worker 'EditData'", verifier)
        self.assertNotIn("Invoke-Worker 'WorkbookPath'", verifier)
        self.assertIn("Invoke-Worker 'MutationUpdate'", verifier)
        self.assertIn("$mutator = Join-Path $scripts 'workbook_probe.py'", verifier)
        self.assertIn("'Primary'", worker)
        self.assertIn("ActivateChartDataWindow", worker)
        self.assertIn("$chartData.Activate()", worker)
        self.assertIn("$presentation.Save()", worker)
        self.assertIn("$activatedWorkbook = $chartData.Workbook", worker)
        self.assertNotIn(
            "GetFullPath([string]$chartData.Workbook.FullName)", worker
        )
        primary_start = worker.index("if ($Mode -eq 'Primary')")
        primary_end = worker.index("} elseif ($Mode -eq 'EditData')", primary_start)
        primary = worker[primary_start:primary_end]
        self.assertLess(
            primary.index("$chartData.Activate()"),
            primary.index("$chartData.ActivateChartDataWindow()"),
        )
        self.assertNotIn("$probeCell.Value2 = $ExpectedValue", primary)
        self.assertNotIn("Primary mode requires ExpectedValue", primary)
        self.assertNotIn("return $fallbackZeroBased + 1", worker)

    def test_office_verification_manifest_uses_real_mapped_point(self):
        source = self.root / "manifest-source.pptx"
        make_fixture(source)
        repair, _ = self._recover(source, "manifest-work")
        built = office_manifest.build_manifest(repair)
        probe = built["charts"][0]["probe"]
        self.assertEqual(built["charts"][0]["chart_ordinal"], 1)
        self.assertEqual(probe["series_name"], "Series 1")
        self.assertEqual(probe["sheet"], "Sheet A")
        self.assertEqual(probe["cell"], "B2")
        self.assertEqual(probe["expected_value"], 10.0)

    def test_office_worker_resolves_duplicate_shape_names_by_chart_ordinal(self):
        worker = (SCRIPTS / "office_stage_worker.ps1").read_text(encoding="utf-8")

        self.assertIn("[int]$expectedOrdinal", worker)
        self.assertIn("$chartInfo.chart_ordinal", worker)
        self.assertIn("$chartShapeIndexes", worker)

    def test_office_verification_manifest_uses_captured_runtime_series_name(self):
        source = self.root / "runtime-series-source.pptx"
        make_fixture(source)
        repair, _ = self._recover(source, "runtime-series-work")
        mapping = repair["charts"][0]["data_mappings"][0]
        mapping["series_name"] = None
        mapping["runtime_series_name"] = "系列1"

        built = office_manifest.build_manifest(repair)

        self.assertEqual(built["charts"][0]["probe"]["series_name"], "系列1")

    def test_office_verification_manifest_rejects_missing_or_none_series_name(self):
        source = self.root / "missing-series-source.pptx"
        make_fixture(source)
        repair, _ = self._recover(source, "missing-series-work")
        mapping = repair["charts"][0]["data_mappings"][0]
        for invalid in (None, "", "None"):
            with self.subTest(invalid=invalid):
                mapping["series_name"] = invalid
                mapping.pop("runtime_series_name", None)
                with self.assertRaisesRegex(
                    office_manifest.VerificationManifestError,
                    "unique runtime series name",
                ):
                    office_manifest.build_manifest(repair)

    def test_office_worker_resolves_combo_series_by_name(self):
        worker = (SCRIPTS / "office_stage_worker.ps1").read_text(encoding="utf-8")
        verifier = (SCRIPTS / "office_verify.ps1").read_text(encoding="utf-8")
        self.assertIn("ExpectedSeriesName", worker)
        self.assertIn("Resolve-SeriesOrdinal", worker)
        self.assertIn("-ExpectedSeriesName", verifier)

    def test_office_gate_ignores_threadless_powerpoint_zombies(self):
        worker = (SCRIPTS / "office_stage_worker.ps1").read_text(encoding="utf-8")
        verifier = (SCRIPTS / "office_verify.ps1").read_text(encoding="utf-8")
        self.assertIn("Get-LivePowerPointPids", worker)
        self.assertIn("Get-LivePowerPointPids", verifier)
        self.assertIn("Threads.Count", worker)
        self.assertIn("Threads.Count", verifier)

    def test_chart_data_activation_never_reuses_user_excel(self):
        worker = (SCRIPTS / "office_stage_worker.ps1").read_text(encoding="utf-8")
        self.assertIn("excelOwnedByWorker", worker)
        self.assertIn("$excelPid -in $pidsBeforeExcel", worker)
        self.assertIn("pre-existing Excel process", worker)

    def test_mutation_worker_tracks_the_unique_excel_started_by_link_update(self):
        worker = (SCRIPTS / "office_stage_worker.ps1").read_text(encoding="utf-8")
        self.assertIn("Resolve-OwnedUpdateExcelPid", worker)
        self.assertIn("Ambiguous Excel processes appeared during LinkFormat.Update", worker)
        self.assertIn("update-excel-pid=", worker)
        self.assertIn("Stop-Process -Id $updateExcelPid", worker)

    def test_office_timeout_scales_with_chart_count(self):
        self.assertGreaterEqual(
            cr.office_verify_timeout_seconds({"charts": [{}] * 8}),
            1200,
        )

    def test_office_cleanup_only_targets_journaled_processes(self):
        worker = (SCRIPTS / "office_stage_worker.ps1").read_text(encoding="utf-8")
        verifier = (SCRIPTS / "office_verify.ps1").read_text(encoding="utf-8")
        portable = (SCRIPTS / "portable_rebind.ps1").read_text(encoding="utf-8")
        carrier = (SCRIPTS / "office_native_carrier.ps1").read_text(encoding="utf-8")
        paste = (SCRIPTS / "office_native_paste.ps1").read_text(encoding="utf-8")
        self.assertIn("powerpoint-pid=", worker)
        self.assertIn("excel-pid=", worker)
        self.assertIn("Stop-TrackedOffice", verifier)
        self.assertIn("ownedOfficePids", portable)
        self.assertIn("Stop-OwnedProcess", carrier)
        self.assertIn("Stop-OwnedProcess", paste)
        self.assertIn("powerpointPid", paste)
        self.assertNotIn("$_.Id -notin $baselinePowerPointPids", verifier)
        self.assertNotIn("$_.Id -notin $baselineExcelPids", verifier)
        self.assertNotIn("$_.Id -notin $baselinePowerPointPids", portable)
        self.assertNotIn("$_.Id -notin $baselineExcelPids", portable)
        self.assertNotIn("$_.Id -notin $baselineExcelPids", carrier)
        self.assertNotIn("$_.Id -notin $baselinePowerPointPids", paste)
        self.assertNotIn("$_.Id -notin $baselineExcelPids", paste)

    def test_office_verifier_is_interrupt_safe_and_propagates_pause(self):
        verifier = (SCRIPTS / "office_verify.ps1").read_text(encoding="utf-8")
        worker = (SCRIPTS / "office_stage_worker.ps1").read_text(encoding="utf-8")
        self.assertIn("if (-not $completeBackup)", verifier)
        self.assertIn("$partialBackup = $backupWorkbook + '.partial'", verifier)
        self.assertIn("$backupHashPath = $backupWorkbook + '.sha256'", verifier)
        self.assertIn("Assert-ValidWorkbookBackup $partialBackup", verifier)
        self.assertIn("Move-Item -LiteralPath $partialBackup -Destination $backupWorkbook", verifier)
        self.assertIn("Copy-Item -LiteralPath $backupWorkbook -Destination $workbookPath", verifier)
        self.assertIn("Remove-Item -LiteralPath $backupWorkbook", verifier)
        self.assertIn("$script:requiresUserAction = $true", verifier)
        self.assertIn("if ($summary.requires_user_action) { exit 20 }", verifier)
        self.assertIn("$powerpointPid -in $pidsBeforePowerPoint", worker)
        self.assertIn("$powerpointPid -notin $pidsBeforePowerPoint", worker)

    def test_native_paste_gate_ignores_threadless_powerpoint_shells(self):
        paste = (SCRIPTS / "office_native_paste.ps1").read_text(encoding="utf-8")
        self.assertIn("Get-LivePowerPointPids", paste)
        self.assertIn("Threads.Count", paste)

    def test_batch_native_paste_uses_one_owned_office_session(self):
        batch_path = SCRIPTS / "office_native_paste_batch.ps1"
        self.assertTrue(batch_path.exists())
        batch = batch_path.read_text(encoding="utf-8")
        self.assertEqual(
            batch.count("New-Object -ComObject PowerPoint.Application"), 1
        )
        self.assertEqual(batch.count("New-Object -ComObject Excel.Application"), 1)
        self.assertIn("foreach ($chart in @($manifest.charts))", batch)
        self.assertIn("runtime_series_names", batch)
        self.assertIn("Stop-OwnedProcess", batch)
        self.assertIn("$maxPasteAttempts = 2", batch)
        self.assertIn("batch-progress.log", batch)
        self.assertNotIn("$newShape.Chart.", batch)
        self.assertIn("Release-Com $linkedChartData", batch)
        self.assertIn("Release-Com $linkedChart", batch)

    def test_final_runtime_series_names_are_captured_after_transplant(self):
        capture_path = SCRIPTS / "office_capture_series_names.ps1"
        self.assertTrue(capture_path.exists())
        capture = capture_path.read_text(encoding="utf-8")
        pipeline = (SCRIPTS / "native_link_pipeline.py").read_text(encoding="utf-8")
        self.assertEqual(
            capture.count("New-Object -ComObject PowerPoint.Application"), 1
        )
        self.assertIn("foreach ($chart in @($manifest.charts))", capture)
        self.assertIn("runtime_series_names", capture)
        self.assertIn("office_capture_series_names.ps1", pipeline)
        self.assertIn("final-series-names.json", pipeline)

    def test_office_verifier_batches_visual_export_and_reports_performance(self):
        export_path = SCRIPTS / "office_export_slides.ps1"
        self.assertTrue(export_path.exists())
        exporter = export_path.read_text(encoding="utf-8")
        verifier = (SCRIPTS / "office_verify.ps1").read_text(encoding="utf-8")
        pipeline = (SCRIPTS / "native_link_pipeline.py").read_text(encoding="utf-8")
        self.assertEqual(
            exporter.count("New-Object -ComObject PowerPoint.Application"), 1
        )
        self.assertIn("Sort-Object -Unique", exporter)
        self.assertIn(".Export(", exporter)
        self.assertIn("office_export_slides.ps1", verifier)
        self.assertIn("visual_exports", verifier)
        self.assertIn("elapsed_seconds", verifier)
        self.assertIn("office_launch_records", verifier)
        self.assertIn("$_.Name -ne 'visual-export-progress.log'", verifier)
        self.assertIn('result["elapsed_seconds"]', pipeline)

    def test_carrier_cleanup_removes_all_temporary_sheets(self):
        source = self.root / "cleanup.pptx"
        make_fixture(source)
        repair, _ = self._recover(source, "cleanup-work")
        workbook = Path(repair["groups"][0]["workbook_path"])
        book = load_workbook(workbook)
        book.create_sheet("L_temp_carrier")
        book.save(workbook)
        book.close()
        recipes = self.root / "recipes.json"
        support = self.root / "support.json"
        recipes.write_text(json.dumps({"charts": repair["charts"]}), encoding="utf-8")
        support.write_text(
            json.dumps(
                [
                    {
                        "workbook_path": str(workbook),
                        "support_sheet": "L_temp_carrier",
                    }
                ]
            ),
            encoding="utf-8",
        )
        result = carriers.clean_support_sheets(recipes, support)
        self.assertTrue(result["passed"])
        reopened = load_workbook(workbook, read_only=True)
        try:
            self.assertNotIn("L_temp_carrier", reopened.sheetnames)
        finally:
            reopened.close()

    def test_portable_bundle_is_self_contained_and_optional(self):
        source = self.root / "portable.pptx"
        make_fixture(source)
        repair, manifest_path = self._recover(source, "portable-work")
        repair["native_relink"] = {"passed": True}
        repair["output_pptx"] = str(source)
        manifest_path.write_text(
            json.dumps(repair, ensure_ascii=False), encoding="utf-8"
        )
        (Path(repair["work_dir"]) / "office-verification.json").write_text(
            json.dumps(
                {
                    "passed": True,
                    "total": 1,
                    "passed_count": 1,
                    "charts": [
                        {
                            "chart_id": repair["charts"][0]["chart_id"],
                            "passed": True,
                            "workbook_restored_exactly": True,
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )
        result = portable_bundle.build_bundle(manifest_path, self.root / "packages")
        bundle = Path(result["bundle_dir"])
        self.assertTrue((bundle / "bundle-manifest.json").exists())
        self.assertTrue((bundle / "本机重绑.ps1").exists())
        self.assertTrue((bundle / "启动并重绑PPT图表.cmd").exists())
        self.assertEqual(len(list((bundle / "data").glob("*.xlsx"))), 1)
        manifest = json.loads(
            (bundle / "bundle-manifest.json").read_text(encoding="utf-8")
        )
        self.assertEqual(manifest["charts"][0]["chart_rel_part"], "ppt/charts/_rels/chart1.xml.rels")
        self.assertEqual(manifest["charts"][0]["probe"]["series_name"], "Series 1")
        self.assertRegex(manifest["bundle_id"], r"^[0-9a-f]{24}$")

    def test_portable_bundle_requires_completed_office_verification(self):
        source = self.root / "portable-unverified.pptx"
        make_fixture(source)
        repair, manifest_path = self._recover(source, "portable-unverified-work")
        repair["native_relink"] = {"passed": True}
        repair["output_pptx"] = str(source)
        manifest_path.write_text(
            json.dumps(repair, ensure_ascii=False), encoding="utf-8"
        )
        with self.assertRaisesRegex(
            portable_bundle.PortableBundleError,
            "Office verification",
        ):
            portable_bundle.build_bundle(manifest_path, self.root / "packages")

    def test_portable_bundle_refuses_to_delete_unmarked_existing_directory(self):
        source = self.root / "portable-existing.pptx"
        make_fixture(source)
        repair, manifest_path = self._recover(source, "portable-existing-work")
        repair["native_relink"] = {"passed": True}
        repair["output_pptx"] = str(source)
        manifest_path.write_text(
            json.dumps(repair, ensure_ascii=False), encoding="utf-8"
        )
        (Path(repair["work_dir"]) / "office-verification.json").write_text(
            json.dumps(
                {
                    "passed": True,
                    "charts": [
                        {
                            "chart_id": repair["charts"][0]["chart_id"],
                            "passed": True,
                            "workbook_restored_exactly": True,
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )
        output_root = self.root / "packages"
        existing = output_root / "portable-existing-图表数据包"
        existing.mkdir(parents=True)
        sentinel = existing / "keep-me.txt"
        sentinel.write_text("user-data", encoding="utf-8")

        with self.assertRaisesRegex(
            portable_bundle.PortableBundleError, "not created by this tool"
        ):
            portable_bundle.build_bundle(manifest_path, output_root)

        self.assertEqual(sentinel.read_text(encoding="utf-8"), "user-data")

    def test_portable_bundle_builds_in_staging_before_replacing_owned_output(self):
        source = self.root / "portable-atomic.pptx"
        make_fixture(source)
        repair, manifest_path = self._recover(source, "portable-atomic-work")
        repair["native_relink"] = {"passed": True}
        repair["output_pptx"] = str(source)
        manifest_path.write_text(
            json.dumps(repair, ensure_ascii=False), encoding="utf-8"
        )
        (Path(repair["work_dir"]) / "office-verification.json").write_text(
            json.dumps(
                {
                    "passed": True,
                    "charts": [
                        {
                            "chart_id": repair["charts"][0]["chart_id"],
                            "passed": True,
                            "workbook_restored_exactly": True,
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )
        output_root = self.root / "packages"
        first = portable_bundle.build_bundle(manifest_path, output_root)
        bundle = Path(first["bundle_dir"])
        original_manifest = (bundle / "bundle-manifest.json").read_bytes()

        with patch(
            "portable_bundle.shutil.copy2",
            side_effect=OSError("simulated staging failure"),
        ):
            with self.assertRaisesRegex(OSError, "simulated staging failure"):
                portable_bundle.build_bundle(manifest_path, output_root)

        self.assertEqual(
            (bundle / "bundle-manifest.json").read_bytes(), original_manifest
        )

    def test_portable_rebinder_uses_safe_runtime_series_and_process_matching(self):
        source = (SCRIPTS / "portable_rebind.ps1").read_text(encoding="utf-8")
        self.assertIn("Get-LivePowerPointPids", source)
        self.assertIn("Resolve-SeriesOrdinal", source)
        self.assertIn("Resolve-OwnedUpdateExcelPid", source)
        self.assertIn("updateExcelPid", source)
        self.assertIn("Ambiguous Excel processes appeared during LinkFormat.Update", source)
        self.assertIn("series_name", source)
        self.assertIn("$excelPid -in $baselineExcelPids", source)
        self.assertIn("数据工作簿哈希校验失败", source)

    def test_office_relink_executes_dependency_probe_before_selecting_python(self):
        source = (SCRIPTS / "office_relink.ps1").read_text(encoding="utf-8")
        self.assertIn("Test-PythonRuntime", source)
        self.assertIn("import lxml, openpyxl", source)
        self.assertIn("WindowsApps", source)


if __name__ == "__main__":
    unittest.main()
