# Copyright 2026 XLoffice-Fyl
# SPDX-License-Identifier: Apache-2.0

"""Make a reversible numeric mutation in a recovered workbook for link testing."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--cell", required=True)
    parser.add_argument("--value", type=float, required=True)
    args = parser.parse_args()
    workbook = args.workbook.resolve()
    wb = load_workbook(workbook)
    try:
        if args.sheet not in wb.sheetnames:
            raise KeyError(f"Worksheet not found: {args.sheet}")
        old_value = wb[args.sheet][args.cell].value
        wb[args.sheet][args.cell] = args.value
        wb.save(workbook)
    finally:
        wb.close()
    print(
        json.dumps(
            {
                "workbook": str(workbook),
                "sheet": args.sheet,
                "cell": args.cell,
                "old_value": old_value,
                "new_value": args.value,
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
