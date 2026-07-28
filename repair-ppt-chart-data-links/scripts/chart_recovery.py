#!/usr/bin/env python3
# Copyright 2026 XLoffice-Fyl
# SPDX-License-Identifier: Apache-2.0

"""Recover PowerPoint chart data links from OOXML caches."""

from __future__ import annotations

import ntpath
import argparse
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import quote, unquote

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook.defined_name import DefinedName

from ooxml_security import (
    OoxmlSecurityError,
    safe_xml_fromstring,
    validate_ooxml_archive,
)


@dataclass(frozen=True)
class FormulaReference:
    workbook: str | None
    sheet: str
    address: str


@dataclass(frozen=True)
class CachePoints:
    values: list[Any]
    declared_count: int
    actual_count: int
    has_gaps: bool
    indexes: list[int]


class ConfirmationRequired(RuntimeError):
    """Raised when a mutating recovery action lacks explicit confirmation."""


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def parse_formula(formula: str) -> FormulaReference:
    if not formula or "!" not in formula:
        raise ValueError(f"Unsupported chart formula: {formula!r}")
    left, address = formula.rsplit("!", 1)
    left = left.strip().strip("'").replace("''", "'")
    workbook = None
    match = re.search(r"\[([^\]]+)\]", left)
    if match:
        workbook = match.group(1)
        sheet = left[match.end() :]
    else:
        sheet = left
    if not sheet:
        raise ValueError(f"Formula has no worksheet name: {formula!r}")
    address = address.replace("$", "")
    return FormulaReference(workbook=workbook, sheet=sheet, address=address)


def cache_points(ref_node: Any, numeric: bool) -> CachePoints:
    cache_name = "numCache" if numeric else "strCache"
    cache = next((node for node in ref_node.iter() if _local_name(node.tag) == cache_name), None)
    if cache is None:
        return CachePoints([], 0, 0, True, [])
    count_node = next((node for node in cache if _local_name(node.tag) == "ptCount"), None)
    points: dict[int, Any] = {}
    for point in cache:
        if _local_name(point.tag) != "pt":
            continue
        idx = int(point.get("idx", "0"))
        value_node = next((node for node in point if _local_name(node.tag) == "v"), None)
        raw = "" if value_node is None or value_node.text is None else value_node.text
        points[idx] = float(raw) if numeric and raw != "" else raw
    declared = int(count_node.get("val")) if count_node is not None else 0
    length = max(declared, max(points, default=-1) + 1)
    values = [points.get(index) for index in range(length)]
    actual = len(points)
    has_gaps = actual != declared or any(value is None for value in values)
    return CachePoints(values, declared, actual, has_gaps, sorted(points))


def multilevel_cache_points(ref_node: Any, address: str) -> CachePoints:
    """Flatten multi-level categories into row-major worksheet cell order."""
    cache = next(
        (node for node in ref_node.iter() if _local_name(node.tag) == "multiLvlStrCache"),
        None,
    )
    if cache is None:
        return CachePoints([], 0, 0, True, [])
    min_col, min_row, max_col, max_row = range_boundaries(address)
    row_count = max_row - min_row + 1
    col_count = max_col - min_col + 1
    levels = [node for node in cache if _local_name(node.tag) == "lvl"]
    count_node = next(
        (node for node in cache if _local_name(node.tag) == "ptCount"), None
    )
    declared_per_level = int(count_node.get("val", "0")) if count_node is not None else 0
    cell_values: dict[tuple[int, int], Any] = {}
    shape_valid = False
    if len(levels) == col_count:
        shape_valid = declared_per_level == row_count
        for level_index, level in enumerate(levels):
            for point in level:
                if _local_name(point.tag) != "pt":
                    continue
                idx = int(point.get("idx", "-1"))
                value_node = next(
                    (node for node in point if _local_name(node.tag) == "v"), None
                )
                if 0 <= idx < row_count and value_node is not None:
                    cell_values[(idx, level_index)] = value_node.text or ""
    elif len(levels) == row_count:
        shape_valid = declared_per_level == col_count
        for level_index, level in enumerate(levels):
            for point in level:
                if _local_name(point.tag) != "pt":
                    continue
                idx = int(point.get("idx", "-1"))
                value_node = next(
                    (node for node in point if _local_name(node.tag) == "v"), None
                )
                if 0 <= idx < col_count and value_node is not None:
                    cell_values[(level_index, idx)] = value_node.text or ""
    values = [
        cell_values.get((row, col))
        for row in range(row_count)
        for col in range(col_count)
    ]
    declared_total = declared_per_level * len(levels)
    return CachePoints(
        values,
        declared_total,
        len(cell_values),
        (not shape_valid) or any(value is None for value in values),
        [
            row * col_count + col
            for row in range(row_count)
            for col in range(col_count)
            if (row, col) in cell_values
        ],
    )


def require_confirmation(manifest: dict[str, Any]) -> None:
    if manifest.get("selection_confirmed") is not True:
        raise ConfirmationRequired(
            "Recovery is blocked until selection_confirmed is explicitly true."
        )


def source_group_key(target: str | None, chart_id: str) -> str:
    if not target:
        return f"chart:{chart_id}"
    decoded = unquote(target).replace("file:///", "", 1).replace("/", "\\")
    return ntpath.normcase(ntpath.normpath(decoded))


def decode_process_bytes(value: bytes | None) -> str:
    if not value:
        return ""
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return value.decode(encoding)
        except UnicodeDecodeError:
            continue
    return value.decode("utf-8", errors="replace")


C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"c": C, "a": A, "p": P, "r": R, "pr": PR}

# Strategy markers used by tests and reports.  A structural externalData patch
# is never enough for full success; Office must rebuild or verify the live link.
NATIVE_RELINK_REQUIRED = "native-office-relink-required"
NATIVE_RELINK_METHOD = (
    "native Excel chart carrier + Office PasteSourceFormatting + OOXML relationship transplant"
)
SUPPORTED_CHART_TYPES = {
    "barChart",
    "lineChart",
    "areaChart",
    "pieChart",
    "doughnutChart",
    "scatterChart",
    "radarChart",
    "bubbleChart",
    "surfaceChart",
    "stockChart",
    "ofPieChart",
}
DATA_ROLES = ("tx", "cat", "val", "xVal", "yVal", "bubbleSize")


class RecoveryBlocked(RuntimeError):
    """Raised when recovery cannot safely proceed."""


def _safe_xml(payload: bytes | str) -> etree._Element:
    try:
        return safe_xml_fromstring(payload)
    except OoxmlSecurityError as exc:
        raise RecoveryBlocked(str(exc)) from exc


def _validate_archive(archive: zipfile.ZipFile) -> None:
    try:
        validate_ooxml_archive(archive)
    except OoxmlSecurityError as exc:
        raise RecoveryBlocked(str(exc)) from exc


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _json_dump(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def _json_load(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _zip_part(base: str, target: str) -> str:
    current = PurePosixPath(base).parent
    joined = current.joinpath(PurePosixPath(target))
    parts: list[str] = []
    for part in joined.parts:
        if part == "..":
            if parts:
                parts.pop()
        elif part not in (".", ""):
            parts.append(part)
    return "/".join(parts)


def _relationships(xml_bytes: bytes) -> dict[str, dict[str, str]]:
    root = _safe_xml(xml_bytes)
    return {
        rel.get("Id"): {
            "target": rel.get("Target", ""),
            "type": rel.get("Type", ""),
            "target_mode": rel.get("TargetMode", ""),
        }
        for rel in root.findall(f"{{{PR}}}Relationship")
    }


def _footer_page(slide_root: etree._Element) -> int | None:
    candidates = []
    for node in slide_root.xpath(".//a:t", namespaces=NS):
        text = (node.text or "").strip()
        if re.fullmatch(r"\d{1,4}", text):
            candidates.append(int(text))
    return candidates[-1] if candidates else None


def _chart_shape_mappings(slide_root: etree._Element) -> list[dict[str, str]]:
    mappings = []
    for frame in slide_root.xpath(".//p:graphicFrame", namespaces=NS):
        chart_nodes = frame.xpath(".//*[@r:id and local-name()='chart']", namespaces=NS)
        if not chart_nodes:
            continue
        chart = chart_nodes[0]
        name_node = frame.find(f".//{{{P}}}cNvPr")
        mappings.append(
            {
                "relationship_id": chart.get(f"{{{R}}}id"),
                "shape_name": name_node.get("name", "") if name_node is not None else "",
                "object_type": "chart" if chart.tag == f"{{{C}}}chart" else "ChartEx",
            }
        )
    return mappings


def _extract_cache(
    ref: etree._Element, address: str | None = None
) -> tuple[str, CachePoints, str | None]:
    if ref.find(f"{{{C}}}numCache") is not None:
        result = cache_points(ref, numeric=True)
        cache_type = "numeric"
    elif ref.find(f"{{{C}}}strCache") is not None:
        result = cache_points(ref, numeric=False)
        cache_type = "string"
    elif ref.find(f"{{{C}}}multiLvlStrCache") is not None:
        result = multilevel_cache_points(ref, address or "A1")
        cache_type = "multilevel"
    else:
        result = CachePoints([], 0, 0, True, [])
        cache_type = "none"
    format_node = ref.find(f".//{{{C}}}formatCode")
    return cache_type, result, format_node.text if format_node is not None else None


def _intentional_gap_kind(cache: CachePoints, expected_count: int) -> str | None:
    """Classify index-preserving edge gaps that PowerPoint intentionally omits."""
    if (
        cache.actual_count == 0
        or len(cache.values) != expected_count
        or any(index >= expected_count for index in cache.indexes)
        or cache.declared_count not in {expected_count, cache.actual_count}
    ):
        return None
    missing = [index for index in range(expected_count) if index not in set(cache.indexes)]
    if not missing:
        return None
    first = min(cache.indexes)
    last = max(cache.indexes)
    leading = list(range(0, first))
    trailing = list(range(last + 1, expected_count))
    if missing == leading:
        return "leading"
    if missing == trailing:
        return "trailing"
    if missing == leading + trailing:
        return "edge"
    return None


def _range_cells(address: str) -> list[str]:
    min_col, min_row, max_col, max_row = range_boundaries(address)
    return [
        f"{get_column_letter(col)}{row}"
        for row in range(min_row, max_row + 1)
        for col in range(min_col, max_col + 1)
    ]


def _analyze_chart(chart_xml: bytes, chart_id: str) -> dict[str, Any]:
    root = _safe_xml(chart_xml)
    local_names = {_local_name(node.tag) for node in root.iter()}
    chart_types = sorted(local_names.intersection(SUPPORTED_CHART_TYPES))
    unsupported = bool(root.find(f".//{{{C}}}pivotSource") is not None)
    references = []
    issues = []
    for series_index, series in enumerate(root.findall(f".//{{{C}}}ser")):
        for role in DATA_ROLES:
            role_node = series.find(f"{{{C}}}{role}")
            if role_node is None:
                continue
            refs = [
                node
                for node in role_node.iter()
                if _local_name(node.tag) in ("strRef", "numRef", "multiLvlStrRef")
            ]
            for ref in refs:
                formula_node = next(
                    (node for node in ref if _local_name(node.tag) == "f"), None
                )
                formula = formula_node.text if formula_node is not None else None
                if not formula:
                    continue
                if "!" not in formula:
                    cache_type, cache, format_code = _extract_cache(ref)
                    incomplete = cache_type == "none" or cache.actual_count == 0
                    if incomplete:
                        issues.append(
                            f"Incomplete defined-name cache for series {series_index} {role}: {formula}"
                        )
                    references.append(
                        {
                            "series_index": series_index,
                            "role": role,
                            "formula": formula,
                            "defined_name": formula.lstrip("="),
                            "workbook": None,
                            "sheet": None,
                            "address": None,
                            "cells": [],
                            "cache_type": cache_type,
                            "values": cache.values,
                            "declared_count": cache.declared_count,
                            "actual_count": cache.actual_count,
                            "has_gaps": cache.has_gaps,
                            "cache_indexes": cache.indexes,
                            "intentional_gap": None,
                            "format_code": format_code,
                        }
                    )
                    continue
                try:
                    parsed = parse_formula(formula)
                    cells = _range_cells(parsed.address)
                except Exception as exc:
                    references.append(
                        {
                            "series_index": series_index,
                            "role": role,
                            "formula": formula,
                            "error": str(exc),
                            "values": [],
                        }
                    )
                    issues.append(f"Unparseable formula: {formula}")
                    continue
                cache_type, cache, format_code = _extract_cache(ref, parsed.address)
                expected_count = len(cells)
                intentional_gap = _intentional_gap_kind(cache, expected_count)
                incomplete = (
                    cache_type == "none"
                    or (cache.has_gaps and intentional_gap is None)
                    or (
                        cache.declared_count != expected_count
                        and intentional_gap is None
                    )
                    or len(cache.values) != expected_count
                )
                if incomplete:
                    issues.append(
                        f"Incomplete cache for series {series_index} {role}: "
                        f"expected {expected_count}, declared {cache.declared_count}, actual {cache.actual_count}"
                    )
                references.append(
                    {
                        "series_index": series_index,
                        "role": role,
                        "formula": formula,
                        "workbook": parsed.workbook,
                        "sheet": parsed.sheet,
                        "address": parsed.address,
                        "cells": cells,
                        "cache_type": cache_type,
                        "values": cache.values,
                        "declared_count": cache.declared_count,
                        "actual_count": cache.actual_count,
                        "has_gaps": cache.has_gaps,
                        "cache_indexes": cache.indexes,
                        "intentional_gap": intentional_gap,
                        "format_code": format_code,
                    }
                )
    if unsupported or not chart_types:
        recoverability = "U"
        issues.append("Unsupported PivotChart, ChartEx, or unknown chart type")
    elif not references:
        recoverability = "C"
        issues.append("No recoverable formula/cache references")
    elif issues:
        recoverability = "B"
    else:
        recoverability = "A"
    title = "".join(node.text or "" for node in root.xpath(".//c:title//a:t", namespaces=NS))
    return {
        "chart_id": chart_id,
        "chart_types": chart_types,
        "references": references,
        "recoverability": recoverability,
        "issues": issues,
        "chart_title": title,
    }


def _detect_conflicts(charts: list[dict[str, Any]]) -> None:
    ledger: dict[tuple[str, str, str], tuple[Any, str]] = {}
    for chart in charts:
        group = source_group_key(chart.get("external_target"), chart["chart_id"])
        for ref in chart.get("references", []):
            values = ref.get("values", [])
            for index, cell in enumerate(ref.get("cells", [])):
                value = values[index] if index < len(values) else None
                if value is None:
                    continue
                key = (group, ref["sheet"], cell)
                typed = (value, ref.get("cache_type", ""))
                if key in ledger and ledger[key] != typed:
                    chart["recoverability"] = "B"
                    chart["issues"].append(
                        f"Conflicting recovered values for {ref['sheet']}!{cell}"
                    )
                else:
                    ledger[key] = typed


def inspect_pptx(pptx: str | Path, work_dir: str | Path) -> dict[str, Any]:
    source = Path(pptx).resolve()
    work = Path(work_dir).resolve()
    if source.suffix.lower() != ".pptx":
        raise RecoveryBlocked("V1 supports .pptx only")
    if not source.exists():
        raise FileNotFoundError(source)
    work.mkdir(parents=True, exist_ok=True)
    charts: list[dict[str, Any]] = []
    text_candidates: list[dict[str, Any]] = []
    with zipfile.ZipFile(source) as archive:
        _validate_archive(archive)
        bad = archive.testzip()
        if bad:
            raise RecoveryBlocked(f"Corrupt ZIP member: {bad}")
        content_types = archive.read("[Content_Types].xml").decode("utf-8", "ignore")
        if "macroEnabled" in content_types:
            raise RecoveryBlocked("Macro-enabled presentations are not supported")
        slide_parts = sorted(
            (
                name
                for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            ),
            key=lambda name: int(re.search(r"slide(\d+)", name).group(1)),
        )
        for slide_part in slide_parts:
            slide_no = int(re.search(r"slide(\d+)", slide_part).group(1))
            slide_root = _safe_xml(archive.read(slide_part))
            footer_page = _footer_page(slide_root)
            texts = [(node.text or "") for node in slide_root.xpath(".//a:t", namespaces=NS)]
            for text in texts:
                if re.search(r"20\d{2}|(?:1\s*[–—-]\s*11)月", text):
                    text_candidates.append(
                        {
                            "slide": slide_no,
                            "footer_page": footer_page,
                            "old_text": text,
                            "suggested_text": None,
                            "reason": "Selected-slide year or period text requires confirmation after data updates",
                            "status": "pending-user-confirmation",
                        }
                    )
            rel_part = f"ppt/slides/_rels/slide{slide_no}.xml.rels"
            rels = _relationships(archive.read(rel_part)) if rel_part in archive.namelist() else {}
            for index, mapping in enumerate(_chart_shape_mappings(slide_root), start=1):
                relationship = rels.get(mapping["relationship_id"], {})
                chart_part = _zip_part(slide_part, relationship.get("target", ""))
                chart_id = f"slide-{slide_no}-chart-{index}"
                if mapping.get("object_type") != "chart":
                    charts.append(
                        {
                            "chart_id": chart_id,
                            "actual_slide": slide_no,
                            "footer_page": footer_page,
                            "shape_name": mapping["shape_name"],
                            "chart_part": chart_part,
                            "recoverability": "U",
                            "issues": ["ChartEx is detected but unsupported in V1"],
                            "selected": False,
                        }
                    )
                    continue
                if chart_part not in archive.namelist():
                    charts.append(
                        {
                            "chart_id": chart_id,
                            "actual_slide": slide_no,
                            "footer_page": footer_page,
                            "shape_name": mapping["shape_name"],
                            "chart_part": chart_part,
                            "recoverability": "U",
                            "issues": ["Chart part missing or ChartEx relationship unsupported"],
                            "selected": False,
                        }
                    )
                    continue
                analysis = _analyze_chart(archive.read(chart_part), chart_id)
                chart_rel_part = f"{PurePosixPath(chart_part).parent}/_rels/{PurePosixPath(chart_part).name}.rels"
                chart_rels = _relationships(archive.read(chart_rel_part)) if chart_rel_part in archive.namelist() else {}
                external_targets = [
                    rel["target"]
                    for rel in chart_rels.values()
                    if rel["target_mode"] == "External"
                ]
                embedded_targets = [
                    rel["target"]
                    for rel in chart_rels.values()
                    if rel["target_mode"] != "External" and rel["type"].endswith("/package")
                ]
                analysis.update(
                    {
                        "actual_slide": slide_no,
                        "footer_page": footer_page,
                        "shape_name": mapping["shape_name"],
                        "chart_part": chart_part,
                        "chart_rel_part": chart_rel_part,
                        "external_target": external_targets[0] if external_targets else None,
                        "embedded_target": embedded_targets[0] if embedded_targets else None,
                        "selected": False,
                    }
                )
                if embedded_targets and not external_targets:
                    analysis["recoverability"] = "N"
                    analysis["issues"] = ["Chart already uses embedded workbook; no external recovery needed"]
                charts.append(analysis)
    _detect_conflicts(charts)
    inventory_path = work / "chart-inventory.json"
    result = {
        "schema_version": "1.0",
        "created_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "source_pptx": str(source),
        "source_sha256": sha256_file(source),
        "work_dir": str(work),
        "selection_confirmed": False,
        "charts": charts,
        "text_candidates": text_candidates,
        "text_candidates_path": str(work / "text-candidates.json"),
        "inventory_path": str(inventory_path),
    }
    _json_dump(work / "text-candidates.json", text_candidates)
    _json_dump(inventory_path, result)
    return result


def confirm_selection(
    inventory_path: str | Path, chart_ids: list[str]
) -> dict[str, Any]:
    path = Path(inventory_path).resolve()
    inventory = _json_load(path)
    requested = list(dict.fromkeys(str(chart_id) for chart_id in chart_ids))
    available = {str(chart["chart_id"]) for chart in inventory.get("charts", [])}
    unknown = sorted(set(requested) - available)
    if unknown:
        raise RecoveryBlocked("Unknown chart IDs: " + ", ".join(unknown))
    if not requested:
        raise RecoveryBlocked("At least one chart ID must be confirmed")
    selected = set(requested)
    for chart in inventory.get("charts", []):
        chart["selected"] = str(chart["chart_id"]) in selected
    inventory["selection_confirmed"] = True
    inventory["confirmed_chart_ids"] = requested
    _json_dump(path, inventory)
    return inventory


def _safe_workbook_name(target: str | None, group_key: str) -> str:
    raw = unquote(target or "")
    basename = ntpath.basename(raw.replace("/", "\\")) or "chart-data.xlsx"
    stem = Path(basename).stem or "chart-data"
    stem = re.sub(r"[<>:\"/\\|?*]", "_", stem).strip(" .") or "chart-data"
    suffix = hashlib.sha256(group_key.encode("utf-8")).hexdigest()[:8]
    return f"{stem}-recovered-{suffix}.xlsx"


def _unique_log_sheet(existing: set[str]) -> str:
    name = "_Recovery_Log"
    index = 2
    while name in existing:
        name = f"_Recovery_Log_{index}"
        index += 1
    return name


def _create_group_workbook(
    output_path: Path,
    charts: list[dict[str, Any]],
    source_hash: str,
) -> tuple[list[dict[str, Any]], list[str]]:
    wb = Workbook()
    wb.remove(wb.active)
    ledger: dict[tuple[str, str], dict[str, Any]] = {}
    named_values: dict[str, dict[str, Any]] = {}
    conflicts: list[str] = []
    for chart in charts:
        for ref in chart.get("references", []):
            values = ref.get("values", [])
            defined_name = ref.get("defined_name")
            if defined_name:
                value = next((value for value in values if value is not None), None)
                previous = named_values.get(str(defined_name))
                if previous is not None and previous["value"] != value:
                    conflicts.append(f"defined-name:{defined_name}")
                else:
                    named_values[str(defined_name)] = {
                        "value": value,
                        "chart_id": chart["chart_id"],
                        "cache_type": ref.get("cache_type"),
                    }
                continue
            for index, cell in enumerate(ref.get("cells", [])):
                value = values[index] if index < len(values) else None
                key = (ref["sheet"], cell)
                entry = {
                    "sheet": ref["sheet"],
                    "cell": cell,
                    "value": value,
                    "cache_type": ref.get("cache_type"),
                    "format_code": ref.get("format_code"),
                    "chart_id": chart["chart_id"],
                    "provenance": "ppt-cache" if value is not None else "missing-cache",
                }
                if key in ledger:
                    previous = ledger[key]
                    if (
                        value is not None
                        and previous["value"] is not None
                        and (value != previous["value"] or entry["cache_type"] != previous["cache_type"])
                    ):
                        conflicts.append(f"{ref['sheet']}!{cell}")
                else:
                    ledger[key] = entry
    sheet_names = {sheet for sheet, _ in ledger}
    if not sheet_names:
        sheet_names = {"RecoveredData"}
    for sheet_name in sorted(sheet_names):
        wb.create_sheet(sheet_name)
    if named_values:
        named_sheet = wb.create_sheet("_NamedRanges")
        named_sheet.append(["Defined name", "Recovered value", "Chart ID"])
        for row, (name, entry) in enumerate(sorted(named_values.items()), start=2):
            named_sheet.cell(row, 1, name)
            named_sheet.cell(row, 2, entry["value"])
            named_sheet.cell(row, 3, entry["chart_id"])
            wb.defined_names.add(
                DefinedName(name, attr_text=f"'_NamedRanges'!$B${row}")
            )
    blue_fill = PatternFill("solid", fgColor="DDEBF7")
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")
    for entry in ledger.values():
        ws = wb[entry["sheet"]]
        cell = ws[entry["cell"]]
        cell.value = entry["value"]
        cell.fill = blue_fill if entry["value"] is not None else yellow_fill
        if entry.get("format_code"):
            cell.number_format = entry["format_code"]
        if entry["value"] is None:
            cell.comment = Comment("Missing from PowerPoint chart cache", "XLoffice-Fyl")
    log_name = _unique_log_sheet(set(wb.sheetnames))
    log = wb.create_sheet(log_name)
    headers = ["Chart ID", "Original target", "Sheet", "Cell", "Value", "Provenance", "Source PPT SHA-256"]
    for col, header in enumerate(headers, start=1):
        cell = log.cell(1, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")
    row = 2
    for entry in sorted(ledger.values(), key=lambda item: (item["sheet"], item["cell"])):
        chart = next(c for c in charts if c["chart_id"] == entry["chart_id"])
        values = [entry["chart_id"], chart.get("external_target"), entry["sheet"], entry["cell"], entry["value"], entry["provenance"], source_hash]
        for col, value in enumerate(values, start=1):
            log.cell(row, col, value)
        row += 1
    for name, entry in sorted(named_values.items()):
        values = [entry["chart_id"], None, "_NamedRanges", name, entry["value"], "ppt-cache-defined-name", source_hash]
        for col, value in enumerate(values, start=1):
            log.cell(row, col, value)
        row += 1
    for col in range(1, len(headers) + 1):
        log.column_dimensions[get_column_letter(col)].width = min(48, max(14, len(headers[col - 1]) + 3))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return list(ledger.values()), sorted(set(conflicts))


def _first_ref(refs: list[dict[str, Any]], role: str) -> dict[str, Any] | None:
    return next((ref for ref in refs if ref.get("role") == role), None)


def build_data_mappings(chart: dict[str, Any]) -> list[dict[str, Any]]:
    mappings: list[dict[str, Any]] = []
    refs_by_series: dict[int, list[dict[str, Any]]] = {}
    for ref in chart.get("references", []):
        refs_by_series.setdefault(int(ref.get("series_index", 0)), []).append(ref)
    for series_index, refs in sorted(refs_by_series.items()):
        tx_ref = _first_ref(refs, "tx")
        cat_ref = _first_ref(refs, "cat")
        for val_ref in [ref for ref in refs if ref.get("role") == "val"]:
            value_cells = list(val_ref.get("cells", []))
            value_values = list(val_ref.get("values", []))
            value_indexes = list(
                val_ref.get(
                    "cache_indexes",
                    [index for index, value in enumerate(value_values) if value is not None],
                )
            )
            category_cells = list(cat_ref.get("cells", [])) if cat_ref else []
            category_values = list(cat_ref.get("values", [])) if cat_ref else []
            formula_point_count = len(value_cells)
            cache_point_count = len(value_indexes)
            mapped_cell_count = sum(
                1
                for index in value_indexes
                if 0 <= index < len(value_cells) and index < len(value_values)
            )
            series_name = None
            if tx_ref and tx_ref.get("values"):
                series_name = tx_ref["values"][0]
            mappings.append(
                {
                    "chart_id": chart["chart_id"],
                    "series_index": series_index,
                    "series_name": series_name,
                    "sheet": val_ref["sheet"],
                    "value_range": val_ref["address"],
                    "value_cells": value_cells,
                    "category_sheet": cat_ref.get("sheet") if cat_ref else None,
                    "category_range": cat_ref.get("address") if cat_ref else None,
                    "category_cells": category_cells,
                    "category_values": category_values,
                    "values": value_values,
                    "cache_indexes": value_indexes,
                    "intentional_gap": val_ref.get("intentional_gap"),
                    "formula_point_count": formula_point_count,
                    "cache_point_count": cache_point_count,
                    "mapped_cell_count": mapped_cell_count,
                    "cache_excel_match": (
                        formula_point_count == len(value_values)
                        and mapped_cell_count == cache_point_count
                        and all(value_values[index] is not None for index in value_indexes)
                    ),
                }
            )
    return mappings


def _bounding_range(cells: list[str]) -> str:
    min_col = min_row = 10**9
    max_col = max_row = 0
    for cell in cells:
        col1, row1, col2, row2 = range_boundaries(cell)
        min_col = min(min_col, col1)
        min_row = min(min_row, row1)
        max_col = max(max_col, col2)
        max_row = max(max_row, row2)
    if max_col == 0:
        return "A1"
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def build_source_data(chart: dict[str, Any]) -> dict[str, Any]:
    sheet_cells: dict[str, list[str]] = {}
    val_shapes = []
    for ref in chart.get("references", []):
        sheet_cells.setdefault(ref["sheet"], []).extend(ref.get("cells", []))
        if ref.get("role") == "val":
            min_col, min_row, max_col, max_row = range_boundaries(ref["address"])
            val_shapes.append((max_col - min_col + 1, max_row - min_row + 1))
    if not sheet_cells:
        return {"sheet": None, "range": None, "formula": None, "plot_by": "xlColumns"}
    if len(sheet_cells) != 1:
        return {"sheet": None, "range": None, "formula": None, "plot_by": "mixed"}
    sheet, cells = next(iter(sheet_cells.items()))
    address = _bounding_range(cells)
    min_col, min_row, max_col, max_row = range_boundaries(address)
    absolute_address = (
        f"${get_column_letter(min_col)}${min_row}:"
        f"${get_column_letter(max_col)}${max_row}"
    )
    row_series = sum(1 for cols, rows in val_shapes if cols > 1 and rows == 1)
    col_series = sum(1 for cols, rows in val_shapes if rows > 1 and cols == 1)
    plot_by = "xlRows" if row_series > col_series else "xlColumns"
    return {
        "sheet": sheet,
        "range": address,
        "formula": f"='{sheet}'!{absolute_address}",
        "plot_by": plot_by,
    }


def recover_manifest(
    inventory_path: str | Path, output_dir: str | Path | None = None
) -> dict[str, Any]:
    path = Path(inventory_path).resolve()
    inventory = _json_load(path)
    require_confirmation(inventory)
    selected = [chart for chart in inventory["charts"] if chart.get("selected")]
    if not selected:
        raise RecoveryBlocked("No charts selected")
    unsupported = [
        chart["chart_id"]
        for chart in selected
        if chart.get("recoverability") in {"C", "U", "N"}
    ]
    if unsupported:
        raise RecoveryBlocked(
            "Selected charts are not eligible for cache recovery: "
            + ", ".join(unsupported)
        )
    for chart in selected:
        immutable = []
        chart["data_mappings"] = build_data_mappings(chart)
        chart["source_data"] = build_source_data(chart)
        for ref in chart.get("references", []):
            if ref.get("error"):
                immutable.append(f"Unparseable formula: {ref.get('formula')}")
            elif (
                not ref.get("defined_name")
                and
                ref.get("declared_count") != len(ref.get("cells", []))
                and ref.get("intentional_gap") is None
            ):
                immutable.append(
                    f"Cache metadata count mismatch for {ref.get('formula')}"
                )
        if any(not mapping["cache_excel_match"] for mapping in chart["data_mappings"]):
            immutable.append("MappedDataMismatch: formula/cache/cell mapping is incomplete")
        chart["immutable_blockers"] = immutable
    groups: dict[str, list[dict[str, Any]]] = {}
    for chart in selected:
        key = source_group_key(chart.get("external_target"), chart["chart_id"])
        groups.setdefault(key, []).append(chart)
    work = Path(inventory["work_dir"])
    delivery_dir = Path(output_dir).resolve() if output_dir is not None else None
    workbook_dir = (
        delivery_dir / "recovered-workbooks"
        if delivery_dir is not None
        else work / "recovered-workbooks"
    )
    group_records = []
    blocked = False
    for key, group_charts in groups.items():
        target = group_charts[0].get("external_target")
        workbook_path = workbook_dir / _safe_workbook_name(target, key)
        ledger, conflicts = _create_group_workbook(
            workbook_path, group_charts, inventory["source_sha256"]
        )
        if conflicts:
            for chart in group_charts:
                chart.setdefault("immutable_blockers", []).extend(
                    f"Conflicting recovered value at {cell}" for cell in conflicts
                )
        if conflicts or any(chart.get("recoverability") != "A" for chart in group_charts):
            blocked = True
        group_records.append(
            {
                "group_key": key,
                "original_target": target,
                "workbook_path": str(workbook_path),
                "chart_ids": [chart["chart_id"] for chart in group_charts],
                "ledger": ledger,
                "conflicts": conflicts,
            }
        )
    manifest_path = work / "repair-manifest.json"
    ppt_output_dir = delivery_dir if delivery_dir is not None else work / "output"
    source = Path(inventory["source_pptx"])
    result = {
        "schema_version": "1.0",
        "created_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "source_pptx": inventory["source_pptx"],
        "source_sha256": inventory["source_sha256"],
        "inventory_path": str(path),
        "work_dir": str(work),
        "selection_confirmed": True,
        "charts": selected,
        "groups": group_records,
        "text_candidates": [
            candidate
            for candidate in inventory.get("text_candidates", [])
            if candidate["slide"] in {chart["actual_slide"] for chart in selected}
        ],
        "text_replacements": [],
        "relink_allowed": not blocked,
        "output_pptx": str(ppt_output_dir / f"{source.stem}-relinked.pptx"),
        "deliverable_report": str(
            (delivery_dir if delivery_dir is not None else work)
            / "数据更新修复报告.md"
        ),
        "relink_strategy": NATIVE_RELINK_REQUIRED,
        "native_relink_method": NATIVE_RELINK_METHOD,
        "manifest_path": str(manifest_path),
    }
    _json_dump(manifest_path, result)
    return result


def _group_for_chart(manifest: dict[str, Any], chart_id: str) -> dict[str, Any]:
    return next(
        group for group in manifest["groups"] if chart_id in group["chart_ids"]
    )


def _values_from_workbook(workbook_path: Path, sheet: str, address: str) -> list[Any]:
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        ws = wb[sheet]
        return [ws[cell].value for cell in _range_cells(address)]
    finally:
        wb.close()


def _recompute_chart_status(manifest: dict[str, Any]) -> None:
    for chart in manifest["charts"]:
        if chart.get("recoverability") == "U":
            continue
        group = _group_for_chart(manifest, chart["chart_id"])
        workbook = Path(group["workbook_path"])
        issues = list(chart.get("immutable_blockers", []))
        for ref in chart.get("references", []):
            values = _values_from_workbook(workbook, ref["sheet"], ref["address"])
            ref["values"] = values
            ref["declared_count"] = len(values)
            ref["actual_count"] = sum(value is not None for value in values)
            ref["has_gaps"] = any(value is None for value in values)
            if ref["has_gaps"]:
                issues.append(f"Missing workbook values for {ref['sheet']}!{ref['address']}")
        chart["issues"] = issues
        chart["recoverability"] = "B" if issues else "A"
    manifest["relink_allowed"] = all(
        chart.get("recoverability") == "A" for chart in manifest["charts"]
    )


def apply_updates(manifest_path: str | Path, updates_path: str | Path) -> dict[str, Any]:
    path = Path(manifest_path).resolve()
    manifest = _json_load(path)
    require_confirmation(manifest)
    payload = _json_load(Path(updates_path).resolve())
    if payload.get("updates_confirmed") is not True:
        raise ConfirmationRequired(
            "Applying user values is blocked until updates_confirmed is explicitly true."
        )
    workbooks: dict[Path, Workbook] = {}
    try:
        for update in payload.get("updates", payload if isinstance(payload, list) else []):
            chart_id = update["chart_id"]
            chart = next(
                (item for item in manifest["charts"] if item["chart_id"] == chart_id),
                None,
            )
            if chart is None:
                raise RecoveryBlocked(f"Update references an unauthorized chart: {chart_id}")
            group = _group_for_chart(manifest, chart_id)
            workbook_path = Path(group["workbook_path"])
            wb = workbooks.setdefault(workbook_path, load_workbook(workbook_path))
            sheet = update.get("sheet")
            if not sheet:
                sheets = {ref["sheet"] for ref in chart.get("references", [])}
                if len(sheets) != 1:
                    raise RecoveryBlocked("Update must specify sheet when chart references multiple sheets")
                sheet = next(iter(sheets))
            if sheet not in wb.sheetnames:
                raise RecoveryBlocked(f"Unknown worksheet for update: {sheet}")
            target_cell = update["target_cell"].replace("$", "").upper()
            allowed = {
                (ref["sheet"], cell.upper())
                for ref in chart.get("references", [])
                for cell in ref.get("cells", [])
            }
            if (sheet, target_cell) not in allowed:
                raise RecoveryBlocked(
                    f"Update cell is outside the selected chart references: {sheet}!{target_cell}"
                )
            if "old_value" not in update:
                raise RecoveryBlocked("Every update must include old_value for audit matching")
            cell = wb[sheet][target_cell]
            if cell.value != update["old_value"]:
                raise RecoveryBlocked(
                    f"Old value mismatch at {sheet}!{target_cell}: "
                    f"expected {update['old_value']!r}, found {cell.value!r}"
                )
            old_value = cell.value
            new_value = update["new_value"]
            old_numeric = isinstance(old_value, (int, float)) and not isinstance(old_value, bool)
            new_numeric = isinstance(new_value, (int, float)) and not isinstance(new_value, bool)
            if old_value is not None and (
                (old_numeric and not new_numeric)
                or (isinstance(old_value, str) and not isinstance(new_value, str))
            ):
                raise RecoveryBlocked(
                    f"New value type does not match cached value at {sheet}!{target_cell}"
                )
            cell.value = new_value
            cell.fill = PatternFill("solid", fgColor="E2F0D9")
            cell.comment = Comment(
                f"User update: {update.get('source_note', 'user supplied')}",
                "XLoffice-Fyl",
            )
            for entry in group["ledger"]:
                if entry["sheet"] == sheet and entry["cell"].upper() == target_cell:
                    entry["value"] = new_value
                    entry["provenance"] = update.get("source_note", "user supplied")
        for workbook_path, wb in workbooks.items():
            wb.save(workbook_path)
    finally:
        for wb in workbooks.values():
            wb.close()
    manifest["text_replacements"] = payload.get("text_replacements", [])
    _recompute_chart_status(manifest)
    _json_dump(path, manifest)
    return manifest


def relink_manifest(manifest_path: str | Path) -> dict[str, Any]:
    """Run the native Excel carrier, Office paste, and OOXML transplant pipeline."""
    from native_link_pipeline import NativeLinkError, RecipeConflict, run_native_relink

    try:
        return run_native_relink(manifest_path)
    except (NativeLinkError, RecipeConflict) as exc:
        raise RecoveryBlocked(str(exc)) from exc


def clean_relink_support_sheets(manifest: dict[str, Any]) -> None:
    """Remove stale linked-chart support sheets from rebuilt workbooks.

    The workbooks handled here are generated by this skill.  Keeping only
    recovered data sheets plus the recovery log makes repeated Office relink
    runs idempotent and prevents Excel worksheet-name collisions.
    """
    for group in manifest.get("groups", []):
        workbook_path = Path(group.get("workbook_path", ""))
        if not workbook_path.exists():
            continue
        keep_sheets = {
            str(entry.get("sheet"))
            for entry in group.get("ledger", [])
            if entry.get("sheet")
        }
        keep_sheets.add("_Recovery_Log")
        if not keep_sheets:
            continue
        wb = load_workbook(workbook_path)
        try:
            for sheet_name in list(wb.sheetnames):
                if sheet_name not in keep_sheets and len(wb.sheetnames) > 1:
                    del wb[sheet_name]
            wb.save(workbook_path)
        finally:
            wb.close()


def _merge_content_types(
    output_content_types: bytes, source_content_types: bytes, restored_parts: set[str]
) -> bytes:
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    output_root = _safe_xml(output_content_types)
    source_root = _safe_xml(source_content_types)
    existing_overrides = {
        node.get("PartName")
        for node in output_root.findall(f"{{{ct_ns}}}Override")
    }
    existing_defaults = {
        node.get("Extension")
        for node in output_root.findall(f"{{{ct_ns}}}Default")
    }
    restored_part_names = {"/" + part for part in restored_parts}
    for node in source_root.findall(f"{{{ct_ns}}}Override"):
        part_name = node.get("PartName")
        if part_name in restored_part_names and part_name not in existing_overrides:
            output_root.append(_safe_xml(etree.tostring(node)))
            existing_overrides.add(part_name)
    for node in source_root.findall(f"{{{ct_ns}}}Default"):
        extension = node.get("Extension")
        if extension not in existing_defaults:
            output_root.append(_safe_xml(etree.tostring(node)))
            existing_defaults.add(extension)
    return etree.tostring(
        output_root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def restore_unselected_ooxml_parts(manifest: dict[str, Any]) -> None:
    """Restore non-target slides and unselected chart parts after Office COM save."""
    if not manifest.get("source_pptx") or not manifest.get("output_pptx"):
        return
    source = Path(manifest["source_pptx"])
    output = Path(manifest["output_pptx"])
    if not source.exists() or not output.exists():
        return
    target_slides = {int(chart["actual_slide"]) for chart in manifest.get("charts", [])}
    native_selected_parts: set[str] = set()
    for chart in manifest.get("charts", []):
        if chart.get("relink_object_kind") == "native-chart":
            if chart.get("chart_part"):
                native_selected_parts.add(chart["chart_part"])
            if chart.get("chart_rel_part"):
                native_selected_parts.add(chart["chart_rel_part"])

    def should_restore(part: str) -> bool:
        slide_match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", part)
        if slide_match:
            return int(slide_match.group(1)) not in target_slides
        slide_rel_match = re.fullmatch(r"ppt/slides/_rels/slide(\d+)\.xml\.rels", part)
        if slide_rel_match:
            return int(slide_rel_match.group(1)) not in target_slides
        if re.fullmatch(r"ppt/charts/chart\d+\.xml", part):
            return part not in native_selected_parts
        if re.fullmatch(r"ppt/charts/_rels/chart\d+\.xml\.rels", part):
            return part not in native_selected_parts
        if re.fullmatch(r"ppt/charts/(?:style|colors)\d+\.xml", part):
            return True
        return False

    fd, temp_name = tempfile.mkstemp(suffix=".pptx", dir=output.parent)
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output) as output_zip:
            _validate_archive(source_zip)
            _validate_archive(output_zip)
            restored_parts = {part for part in source_zip.namelist() if should_restore(part)}
            with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as new_zip:
                written: set[str] = set()
                for item in output_zip.infolist():
                    data = output_zip.read(item.filename)
                    if item.filename in restored_parts:
                        data = source_zip.read(item.filename)
                    if item.filename == "[Content_Types].xml":
                        data = _merge_content_types(
                            data,
                            source_zip.read("[Content_Types].xml"),
                            restored_parts,
                        )
                    new_zip.writestr(item, data)
                    written.add(item.filename)
                for part in sorted(restored_parts - written):
                    new_zip.writestr(part, source_zip.read(part))
        last_error: OSError | None = None
        for _ in range(20):
            try:
                os.replace(temp_path, output)
                last_error = None
                break
            except PermissionError as exc:
                last_error = exc
                time.sleep(0.5)
        if last_error is not None:
            raise last_error
    finally:
        if temp_path.exists():
            temp_path.unlink()


def office_relink_manifest(manifest_path: str | Path) -> dict[str, Any]:
    """Compatibility entry point for the single supported native relink engine."""
    return relink_manifest(manifest_path)


def _target_path(target: str) -> Path:
    decoded = unquote(target)
    if decoded.startswith("file:///"):
        decoded = decoded[8:]
    return Path(decoded.replace("/", os.sep))


def _file_uri(path: Path) -> str:
    """Return the legacy file URI spelling emitted by PowerPoint for Windows."""
    escaped: list[str] = []
    for character in str(path.resolve()):
        if character == " ":
            escaped.append("%20")
        elif character == "#":
            escaped.append("%23")
        elif character == "%":
            escaped.append("%25")
        elif character == "?":
            escaped.append("%3F")
        elif ord(character) < 0x20:
            escaped.append(f"%{ord(character):02X}")
        else:
            escaped.append(character)
    return "file:///" + "".join(escaped)


def _external_data_relationship_id(chart_xml: bytes) -> str:
    root = _safe_xml(chart_xml)
    external_data = root.find(f".//{{{C}}}externalData")
    relationship_id = (
        external_data.get(f"{{{R}}}id") if external_data is not None else None
    )
    if not relationship_id:
        raise RecoveryBlocked("Selected chart has no externalData relationship id")
    return relationship_id


def _chart_cache_value(value: Any, numeric: bool, epoch: dt.datetime) -> Any:
    if numeric and isinstance(value, (dt.datetime, dt.date, dt.time)):
        return to_excel(value, epoch)
    return value


def _verify_structural(manifest: dict[str, Any]) -> tuple[list[dict[str, Any]], bool]:
    checks: list[dict[str, Any]] = []

    def add(name: str, passed: bool, detail: str = "") -> None:
        checks.append({"name": name, "passed": bool(passed), "detail": detail})

    source = Path(manifest["source_pptx"])
    output = Path(manifest["output_pptx"])
    add("source hash unchanged", sha256_file(source) == manifest["source_sha256"])
    add("output exists", output.exists())
    if not output.exists():
        return checks, False
    try:
        with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output) as output_zip:
            _validate_archive(source_zip)
            _validate_archive(output_zip)
            add("output ZIP valid", output_zip.testzip() is None)
            source_slides = {
                name
                for name in source_zip.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            }
            output_slides = {
                name
                for name in output_zip.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            }
            source_charts = {
                name
                for name in source_zip.namelist()
                if re.fullmatch(r"ppt/charts/chart\d+\.xml", name)
            }
            output_charts = {
                name
                for name in output_zip.namelist()
                if re.fullmatch(r"ppt/charts/chart\d+\.xml", name)
            }
            add("slide parts unchanged", source_slides == output_slides)
            add("chart parts unchanged", source_charts == output_charts)
            style_color_parts = {
                name
                for name in source_zip.namelist()
                if re.fullmatch(r"ppt/charts/(?:style|colors)\d+\.xml", name)
            }
            add(
                "all chart style/color parts unchanged",
                all(
                    part in output_zip.namelist()
                    and source_zip.read(part) == output_zip.read(part)
                    for part in style_color_parts
                ),
            )
            selected_parts = {
                chart["chart_rel_part"]
                for chart in manifest["charts"]
                if chart.get("chart_rel_part")
            }
            all_source_chart_rels = {
                name
                for name in source_zip.namelist()
                if name.startswith("ppt/charts/_rels/chart") and name.endswith(".xml.rels")
            }
            add(
                "unselected chart relationships unchanged",
                all(
                    source_zip.read(name) == output_zip.read(name)
                    for name in all_source_chart_rels - selected_parts
                ),
            )
            for chart in manifest["charts"]:
                group = _group_for_chart(manifest, chart["chart_id"])
                workbook = Path(group["workbook_path"]).resolve()
                chart_xml = output_zip.read(chart["chart_part"])
                relationship_id = _external_data_relationship_id(chart_xml)
                rels = _relationships(output_zip.read(chart["chart_rel_part"]))
                target_rel = rels.get(relationship_id)
                targets = [target_rel["target"]] if target_rel else []
                add(
                    f"{chart['chart_id']} target",
                    len(targets) == 1 and _target_path(targets[0]).resolve() == workbook,
                    targets[0] if targets else "missing",
                )
                root = _safe_xml(chart_xml)
                # Random access against openpyxl read-only worksheets repeatedly
                # reparses XML and becomes quadratic on long chart series.  These
                # recovered workbooks are small enough to load normally, making
                # full-range cache equality practical instead of sampling.
                wb = load_workbook(workbook, read_only=False, data_only=False)
                try:
                    mismatches = []
                    for ref in root.xpath(
                        ".//c:strRef | .//c:numRef | .//c:multiLvlStrRef",
                        namespaces=NS,
                    ):
                        formula_node = ref.find(f"{{{C}}}f")
                        if formula_node is None or not formula_node.text:
                            continue
                        parsed = parse_formula(formula_node.text)
                        numeric = _local_name(ref.tag) == "numRef"
                        expected = [
                            _chart_cache_value(
                                wb[parsed.sheet][cell].value, numeric, wb.epoch
                            )
                            for cell in _range_cells(parsed.address)
                        ]
                        if _local_name(ref.tag) == "multiLvlStrRef":
                            actual = multilevel_cache_points(ref, parsed.address).values
                        else:
                            actual = cache_points(
                                ref, numeric=_local_name(ref.tag) == "numRef"
                            ).values
                        if len(expected) != len(actual):
                            mismatches.append(formula_node.text)
                            continue
                        for exp, act in zip(expected, actual):
                            if isinstance(exp, (int, float)) and isinstance(act, (int, float)):
                                if abs(float(exp) - float(act)) > 1e-9:
                                    mismatches.append(formula_node.text)
                                    break
                            elif str(exp) != str(act):
                                mismatches.append(formula_node.text)
                                break
                    add(f"{chart['chart_id']} workbook/cache equality", not mismatches, "; ".join(mismatches))
                finally:
                    wb.close()
    except Exception as exc:
        add("structural verification exception", False, repr(exc))
    return checks, all(check["passed"] for check in checks)


def _write_report(manifest: dict[str, Any], checks: list[dict[str, Any]], office: Any) -> Path:
    work = Path(manifest["work_dir"])
    report = Path(manifest.get("deliverable_report", work / "validation-report.md"))
    report.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# PPT chart link recovery validation report",
        "",
        f"- Source: `{manifest['source_pptx']}`",
        f"- Output: `{manifest['output_pptx']}`",
        f"- Source SHA-256: `{manifest['source_sha256']}`",
        "",
        "| Check | Result | Detail |",
        "|---|---|---|",
    ]
    for check in checks:
        lines.append(
            f"| {check['name']} | {'PASS' if check['passed'] else 'FAIL'} | {check.get('detail', '')} |"
        )
    if office is not None:
        lines.extend(["", "## Office verification", "", f"```json\n{json.dumps(office, ensure_ascii=False, indent=2)}\n```"])
    report.write_text("\n".join(lines), encoding="utf-8")
    return report


def office_verify_timeout_seconds(manifest: dict[str, Any]) -> int:
    """Allow isolated Office stages enough time to verify every selected chart."""
    return max(600, 180 * max(1, len(manifest.get("charts", []))))


def verify_manifest(manifest_path: str | Path, office_required: bool = False) -> dict[str, Any]:
    path = Path(manifest_path).resolve()
    manifest = _json_load(path)
    checks, passed = _verify_structural(manifest)
    office_result = None
    if office_required and passed:
        script = Path(__file__).with_name("office_verify.ps1")
        command = [
            "powershell",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script),
            "-ManifestPath",
            str(path),
            "-PythonPath",
            sys.executable,
        ]
        completed = subprocess.run(
            command,
            capture_output=True,
            text=False,
            timeout=office_verify_timeout_seconds(manifest),
        )
        stdout_text = decode_process_bytes(completed.stdout)
        stderr_text = decode_process_bytes(completed.stderr)
        office_json = Path(manifest["work_dir"]) / "office-verification.json"
        office_result = _json_load(office_json) if office_json.exists() else {"passed": False, "stderr": stderr_text, "stdout": stdout_text}
        if completed.returncode == 20 and office_result.get("requires_user_action"):
            checks.append(
                {
                    "name": "Office verification",
                    "passed": False,
                    "detail": office_result.get("message", "PowerPoint must be closed"),
                }
            )
            passed = False
        else:
            checks.append({"name": "Office verification", "passed": completed.returncode == 0 and office_result.get("passed") is True, "detail": stderr_text.strip()})
            post_checks, post_passed = _verify_structural(manifest)
            checks.extend({**check, "name": f"post-Office: {check['name']}"} for check in post_checks)
            passed = passed and post_passed and checks[-len(post_checks)-1]["passed"]
    report = _write_report(manifest, checks, office_result)
    result = {
        "passed": passed and all(check["passed"] for check in checks),
        "requires_user_action": bool(
            office_result and office_result.get("requires_user_action")
        ),
        "message": office_result.get("message") if office_result else None,
        "checks": checks,
        "office": office_result,
        "report_path": str(report),
    }
    _json_dump(Path(manifest["work_dir"]) / "validation-report.json", result)
    return result


def _cmd_inspect(args: argparse.Namespace) -> int:
    result = inspect_pptx(args.pptx, args.work_dir)
    print(result["inventory_path"])
    return 0


def _cmd_recover(args: argparse.Namespace) -> int:
    result = recover_manifest(args.manifest, args.output_dir)
    print(result["manifest_path"])
    return 0


def _cmd_confirm_selection(args: argparse.Namespace) -> int:
    result = confirm_selection(args.manifest, args.chart_id)
    print(json.dumps(result["confirmed_chart_ids"], ensure_ascii=False))
    return 0


def _cmd_apply_updates(args: argparse.Namespace) -> int:
    result = apply_updates(args.manifest, args.updates)
    print(result["manifest_path"])
    return 0


def _cmd_relink(args: argparse.Namespace) -> int:
    result = relink_manifest(args.manifest)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 20 if result.get("requires_user_action") else 0


def _cmd_office_relink(args: argparse.Namespace) -> int:
    result = office_relink_manifest(args.manifest)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 20 if result.get("requires_user_action") else 0


def _cmd_verify(args: argparse.Namespace) -> int:
    result = verify_manifest(args.manifest, office_required=args.office_required)
    print(result["report_path"])
    if result.get("requires_user_action"):
        print(result.get("message", "PowerPoint must be closed"), file=sys.stderr)
        return 20
    return 0 if result["passed"] else 1


def _cmd_portable_package(args: argparse.Namespace) -> int:
    from portable_bundle import PortableBundleError, build_bundle

    try:
        result = build_bundle(
            Path(args.manifest), Path(args.output_dir) if args.output_dir else None
        )
    except PortableBundleError as exc:
        raise RecoveryBlocked(str(exc)) from exc
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    inspect_parser = subparsers.add_parser("inspect")
    inspect_parser.add_argument("--pptx", required=True)
    inspect_parser.add_argument("--work-dir", required=True)
    inspect_parser.set_defaults(func=_cmd_inspect)
    recover_parser = subparsers.add_parser("recover")
    recover_parser.add_argument("--manifest", required=True)
    recover_parser.add_argument("--output-dir")
    recover_parser.set_defaults(func=_cmd_recover)
    selection_parser = subparsers.add_parser("confirm-selection")
    selection_parser.add_argument("--manifest", required=True)
    selection_parser.add_argument("--chart-id", action="append", required=True)
    selection_parser.set_defaults(func=_cmd_confirm_selection)
    updates_parser = subparsers.add_parser("apply-updates")
    updates_parser.add_argument("--manifest", required=True)
    updates_parser.add_argument("--updates", required=True)
    updates_parser.set_defaults(func=_cmd_apply_updates)
    relink_parser = subparsers.add_parser("relink")
    relink_parser.add_argument("--manifest", required=True)
    relink_parser.set_defaults(func=_cmd_relink)
    office_relink_parser = subparsers.add_parser("office-relink")
    office_relink_parser.add_argument("--manifest", required=True)
    office_relink_parser.set_defaults(func=_cmd_office_relink)
    verify_parser = subparsers.add_parser("verify")
    verify_parser.add_argument("--manifest", required=True)
    verify_parser.add_argument("--office-required", action="store_true")
    verify_parser.set_defaults(func=_cmd_verify)
    portable_parser = subparsers.add_parser("portable-package")
    portable_parser.add_argument("--manifest", required=True)
    portable_parser.add_argument("--output-dir")
    portable_parser.set_defaults(func=_cmd_portable_package)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    try:
        return args.func(args)
    except (ConfirmationRequired, RecoveryBlocked, FileNotFoundError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
